import io
import os
from typing import Union, BinaryIO
from pdfminer.high_level import extract_text as extract_pdf_text
import docx

def extract_text(file_source: Union[str, BinaryIO, io.BytesIO], file_extension: str) -> str:
    """
    Extracts raw text from a PDF or DOCX file.
    
    Args:
        file_source: A file path string, or a file-like object containing bytes.
        file_extension: The file extension ('.pdf' or '.docx'), case-insensitive.
        
    Returns:
        The extracted raw text as a string.
        
    Raises:
        ValueError: If the file extension is unsupported or empty.
        Exception: If extraction fails.
    """
    ext = file_extension.lower().strip()
    if not ext.startswith('.'):
        ext = '.' + ext
        
    if ext == '.pdf':
        return extract_text_from_pdf(file_source)
    elif ext == '.docx':
        return extract_text_from_docx(file_source)
    else:
        raise ValueError(f"Unsupported file format: {ext}. Only PDF and DOCX are supported.")

def extract_text_from_pdf(file_source: Union[str, BinaryIO, io.BytesIO]) -> str:
    """Extracts text from a PDF file source using pdfminer.six."""
    try:
        # If it's bytes, wrap it in BytesIO
        if isinstance(file_source, bytes):
            file_source = io.BytesIO(file_source)
            
        text = extract_pdf_text(file_source)
        if not text:
            return ""
        return text.strip()
    except Exception as e:
        raise Exception(f"Failed to extract text from PDF: {str(e)}")

def extract_text_from_docx(file_source: Union[str, BinaryIO, io.BytesIO]) -> str:
    """Extracts text from a DOCX file source using python-docx."""
    try:
        if isinstance(file_source, bytes):
            file_source = io.BytesIO(file_source)
            
        doc = docx.Document(file_source)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
            
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text.append(cell.text)
                    
        return "\n".join(full_text).strip()
    except Exception as e:
        raise Exception(f"Failed to extract text from DOCX: {str(e)}")
